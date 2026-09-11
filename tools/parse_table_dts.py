import openpyxl
import json
import re
from pathlib import Path

SRC = Path(r"C:\Users\heliu\Desktop\to Huang Pharma-DTS-Std limitation 26097.xlsx")
OUT = Path(r"C:\Users\heliu\Desktop\SCHEM.EXE\data\dts_db.json")

wb = openpyxl.load_workbook(SRC, data_only=True)
ws = wb.worksheets[0]

models = []
for c in range(3, 14):
    value = ws.cell(3, c).value
    if isinstance(value, str) and value.strip():
        models.append((c, value.strip()))

records = []
for r in range(4, 9):
    if ws.cell(r, 1).value != "IC":
        continue
    temp_raw = ws.cell(r, 2).value or ""
    tm = re.search(r"(\d+)\s*~\s*(\d+)", str(temp_raw))
    if not tm:
        continue
    material_in = int(tm.group(1))
    material_out = int(tm.group(2))

    for col, model in models:
        cell = ws.cell(r, col).value
        if not isinstance(cell, str):
            continue
        valid_lines = []
        for line in cell.splitlines():
            line = line.strip()
            if not line:
                continue
            m = re.match(r"^\s*([\d,]+)\s*[（(]([^）)]*)[）)]\s*([\d,]+)", line)
            if not m:
                continue
            material_flow = float(m.group(1).replace(",", ""))
            inner = m.group(2)
            service_flow = float(m.group(3).replace(",", ""))
            margin_m = re.search(r"(\d+(?:\.\d+)?)\s*余量", inner)
            margin = float(margin_m.group(1)) if margin_m else 10.0
            if "0余量" in inner or "余量0" in inner:
                margin = 0.0
            valid_lines.append({
                "material_flow": material_flow,
                "service_flow": service_flow,
                "margin": margin,
                "raw": line
            })

        conditions = ["7-12℃10%_margin", "50kpa", "100kpa", "100kpa_no_margin"]
        pressure_values = [None, 50, 100, 100]
        for idx, line in enumerate(valid_lines[:4]):
            records.append({
                "model": model,
                "material_in": material_in,
                "material_out": material_out,
                "material_flow": line["material_flow"],
                "service_flow": line["service_flow"],
                "margin": line["margin"],
                "pressure": pressure_values[idx],
                "condition": conditions[idx],
                "raw": line["raw"]
            })

OUT.parent.mkdir(parents=True, exist_ok=True)
OUT.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")
print(f"dts_records={len(records)} dts_models={len(models)}")
